#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Generate HTML preview from Word template filled with database record.
Uses Aspose.Words for HTML conversion to preserve exact template styling.

Usage:
  python generate_preview.py <template.docx> <record_json_file> <app_dir>

Output: HTML string to stdout (UTF-8)
"""
import sys
import os
import json
import subprocess

# Force UTF-8 stdout/stderr on Windows
if sys.platform == 'win32':
    try:
        sys.stdout.reconfigure(encoding='utf-8')
        sys.stderr.reconfigure(encoding='utf-8')
    except Exception:
        pass

from docx import Document


def find_content_cell(table, row_idx, target_merged_label_left, offset_from_label=1):
    row = table.rows[row_idx]
    matches = []
    for i, cell in enumerate(row.cells):
        if cell.text.strip() == target_merged_label_left:
            matches.append(i)
    if not matches:
        return None
    label_last_col = matches[-1]
    content_start_col = label_last_col + 1
    if content_start_col >= len(row.cells):
        return None
    return row.cells[content_start_col]


def find_span_cell_by_first_label(table, row_idx, first_label):
    row = table.rows[row_idx]
    matches = [i for i, c in enumerate(row.cells) if c.text.strip() == first_label]
    if matches:
        start = matches[-1] + 1
    else:
        start = 1
    for i in range(start, len(row.cells)):
        cell = row.cells[i]
        if cell.text.strip() == '':
            return cell
    return row.cells[-1] if row.cells else None


def set_cell_text(cell, text):
    if cell is None:
        return
    text = text or ''
    if cell.paragraphs:
        first_para = cell.paragraphs[0]
        for run in first_para.runs:
            run.text = ''
        if first_para.runs:
            first_para.runs[0].text = text
        else:
            first_para.add_run(text)
    else:
        cell.add_paragraph(text)


def fill_template(template_path, record, output_docx_path):
    doc = Document(template_path)
    if not doc.tables:
        raise RuntimeError("模板中找不到表格")
    table = doc.tables[0]

    department = record.get('DEPARTMENT') or ''
    word_code = record.get('WORD_CODE') or ''
    receive_date = record.get('FILE_RECEIVE_DATE') or ''
    file_category = record.get('FILE_CATEGORY') or ''
    receive_channel = record.get('RECEIVE_CHANNEL') or ''
    emergency = record.get('EMERGENCY_LEVEL') or ''
    secret = record.get('SECRET_LEVEL') or ''
    receive_number = record.get('RECEIVE_NUMBER') or ''
    summary = record.get('SUMMARY') or ''
    leader = record.get('LEADER_INSTRUCTION') or ''
    suggestion = record.get('SUGGESTION') or ''
    process_result = record.get('PROCESS_RESULT') or ''

    c = find_content_cell(table, 0, '来文单位')
    set_cell_text(c, department)

    c = find_content_cell(table, 1, '来文字号')
    set_cell_text(c, word_code)
    c = find_content_cell(table, 1, '收文日期')
    set_cell_text(c, receive_date)

    c = find_content_cell(table, 2, '来文类型')
    set_cell_text(c, file_category)
    c = find_content_cell(table, 2, '收文途径')
    set_cell_text(c, receive_channel)

    c = find_content_cell(table, 3, '紧急程度')
    set_cell_text(c, emergency)
    c = find_content_cell(table, 3, '密 级')
    if c is None:
        c = find_content_cell(table, 3, '密级')
    set_cell_text(c, secret)
    c = find_content_cell(table, 3, '收文编号')
    set_cell_text(c, receive_number)

    c = find_span_cell_by_first_label(table, 4, '文件标题')
    set_cell_text(c, summary)

    c = find_span_cell_by_first_label(table, 5, '领导批示')
    set_cell_text(c, leader)

    c = find_span_cell_by_first_label(table, 6, '拟办意见')
    set_cell_text(c, suggestion)

    c = find_span_cell_by_first_label(table, 9, '办理结果')
    set_cell_text(c, process_result)

    doc.save(output_docx_path)


def find_license(app_dir):
    candidates = [
        os.path.join(app_dir, '1.lic'),
        os.path.join(app_dir, 'aspose.words.lic'),
    ]
    for p in candidates:
        if os.path.exists(p):
            return p
    return ''


def docx_to_html(docx_path, app_dir):
    """Convert docx to HTML using Aspose.Words, return HTML string."""
    import aspose.words as aw

    lic = find_license(app_dir)
    if lic:
        try:
            license_obj = aw.License()
            license_obj.set_license(lic)
        except Exception:
            pass

    doc = aw.Document(docx_path)

    import io
    stream = io.BytesIO()
    doc.save(stream, aw.SaveFormat.HTML)
    html = stream.getvalue().decode('utf-8', errors='replace')
    return html


def main():
    if len(sys.argv) < 4:
        print("Usage: generate_preview.py <template.docx> <record_json_file> <app_dir>")
        sys.exit(1)

    template_path = sys.argv[1]
    record_json_path = sys.argv[2]
    app_dir = sys.argv[3]

    with open(record_json_path, 'r', encoding='utf-8') as f:
        record = json.load(f)

    import tempfile
    with tempfile.NamedTemporaryFile(suffix='.docx', delete=False) as tmp:
        tmp_docx = tmp.name

    try:
        fill_template(template_path, record, tmp_docx)
        html = docx_to_html(tmp_docx, app_dir)
        print(html)
    finally:
        try:
            os.unlink(tmp_docx)
        except Exception:
            pass


if __name__ == '__main__':
    main()
